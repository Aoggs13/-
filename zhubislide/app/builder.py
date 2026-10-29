"""Deck JSON + photos → designed 16:9 PPTX (python-pptx).

Slides carry a lot of text (full sentences), so every text block is sized by `fit()`: it simulates word wrap
for the box and picks the largest font size that fills the box without overflowing.
"""

from __future__ import annotations

import contextvars
import datetime as dt
from pathlib import Path

from lxml import etree
from PIL import Image, ImageEnhance, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from .images import fit as fit_image
from .themes import FONTS, STYLES, contrast, resolve, resolve_style

W, H = Inches(13.333), Inches(7.5)
IMAGE_LAYOUTS = {"image_right", "image_left", "quote"}
X0, X1 = 0.75, 12.58     # content margins, inches
TOP, BOTTOM = 2.0, 6.75  # content area under the heading
LINE_H = 1.22            # line height (em) at line_spacing 1.0
SAFETY = 1.1             # glyphs a bit wider than measured: the font may be substituted (Google Slides, phones)

# per-build font and style (builds run in worker threads, ContextVars keep them independent)
_font: contextvars.ContextVar[str] = contextvars.ContextVar("font", default="Arial")
_style: contextvars.ContextVar[dict] = contextvars.ContextVar("style", default={"key": "modern", **STYLES["modern"]})


def font() -> str:
    return _font.get()


def style() -> dict:
    return _style.get()


def needs_image(s: dict) -> bool:
    return s.get("layout") in IMAGE_LAYOUTS


# ───────────────────────── text fitting ─────────────────────────

def _char_w(bold: bool) -> float:
    return (FONTS.get(font(), FONTS["Arial"])["char_w"] + (0.05 if bold else 0)) * SAFETY


def n_lines(value: str, width_in: float, size: float, bold: bool = False) -> int:
    char_w = _char_w(bold)
    per_line = max(4, int(width_in / (size * char_w / 72)))
    total = 0
    for para in value.split("\n"):
        cur, n = 0, 1
        for word in para.split():
            ln = len(word)
            if cur == 0:
                cur = ln
            elif cur + 1 + ln <= per_line:
                cur += 1 + ln
            else:
                n, cur = n + 1, ln
            while cur > per_line:
                n, cur = n + 1, cur - per_line
        total += n
    return total


def text_height(paras: list[str], width_in: float, size: float, spacing: float = 1.1, gap: float = 0.5,
                indent: float = 0.0, bold: bool = False, first_scale: float = 1.0) -> float:
    h = 0.0
    for i, p in enumerate(paras):
        s = size * (first_scale if i == 0 else 1.0)
        ind = 0.0 if (i == 0 and first_scale != 1.0) else indent * s / 72
        h += n_lines(p, width_in - ind, s, bold) * s * LINE_H * spacing / 72
        if i < len(paras) - 1:
            h += s * gap / 72
    return h


def word_fits(paras: list[str], width_in: float, size: float, bold: bool = False) -> bool:
    """The longest word must fit on one line, otherwise it would be broken in the middle."""
    longest = max((len(w) for p in paras for w in p.split()), default=0)
    return longest * size * _char_w(bold) / 72 <= width_in


def fit(paras: list[str], width_in: float, height_in: float, max_pt: float, min_pt: float = 10.0,
        max_lines: int = 0, **kw) -> float:
    """Largest size (max_pt → min_pt) at which the text fits the box, no word is broken and, when given,
    the text takes at most max_lines lines."""
    paras = [p for p in paras if p]
    if not paras:
        return max_pt
    size = max_pt
    bold = kw.get("bold", False)

    def too_big(sz: float) -> bool:
        if text_height(paras, width_in, sz, **kw) > height_in or not word_fits(paras, width_in, sz, bold):
            return True
        return bool(max_lines) and sum(n_lines(p, width_in, sz, bold) for p in paras) > max_lines

    while size > min_pt and too_big(size):
        size -= 0.5
    return size


# ───────────────────────── low-level helpers ─────────────────────────

def I(v: float) -> Emu:  # noqa: E743 - inches; short because the layouts use it everywhere
    return Inches(v)


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


def on(color: str) -> str:
    """Text colour for a label sitting on a filled accent shape."""
    return "FFFFFF" if contrast("FFFFFF", color) >= contrast("111827", color) else "111827"


def shade(t: dict) -> str:
    """Overlay colour on top of photos: always dark, so white text stays readable on light themes too."""
    return t["bg"][0] if t["dark"] else "0F172A"


def _alpha(shape, opacity: float) -> None:
    clr = shape.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    etree.SubElement(clr, qn("a:alpha")).set("val", str(int(opacity * 100000)))


def rect(slide, x, y, w, h, color: str, opacity: float = 1.0, shape=MSO_SHAPE.RECTANGLE, radius: float | None = None):
    sh = slide.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(color)
    sh.line.fill.background()
    sh.shadow.inherit = False
    if radius is not None:
        sh.adjustments[0] = radius
    if opacity < 1:
        _alpha(sh, opacity)
    return sh


def outline(slide, x, y, w, h, color: str, width: float = 1.0, shape=MSO_SHAPE.RECTANGLE):
    sh = slide.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    sh.fill.background()
    sh.shadow.inherit = False
    sh.line.color.rgb = rgb(color)
    sh.line.width = Pt(width)
    return sh


def panel(slide, t, x, y, w, h, radius=0.06, border: str | None = None):
    """Card behind text blocks, drawn in the deck's style."""
    kind = style()["card"]
    fill_op = 0.88 if t["dark"] else 0.95
    if kind == "outline":
        return outline(slide, x, y, w, h, border or t["muted"])
    if kind == "flat":
        return rect(slide, x, y, w, h, t["surface"], fill_op)
    if kind == "bubble":
        p = rect(slide, x, y, w, h, t["surface"], fill_op, MSO_SHAPE.ROUNDED_RECTANGLE, 0.12)
        p.line.color.rgb = rgb(border or t["accent"])
        p.line.width = Pt(2.25)
        return p
    p = rect(slide, x, y, w, h, t["surface"], fill_op, MSO_SHAPE.ROUNDED_RECTANGLE, radius)
    if border:
        p.line.color.rgb = rgb(border)
        p.line.width = Pt(0.75)
    return p


def band_shape():
    """Shape for filled bars (two-column headers): square in flat/outline styles."""
    return MSO_SHAPE.RECTANGLE if style()["card"] in ("flat", "outline") else MSO_SHAPE.ROUNDED_RECTANGLE


def _frame(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def _run(p, value: str, size: float, color: str, bold=False, italic=False):
    r = p.add_run()
    r.text = value
    f = r.font
    f.name, f.size, f.bold, f.italic = font(), Pt(size), bold, italic
    f.color.rgb = rgb(color)
    return r


def text(slide, x, y, w, h, value: str, size: float, color: str, bold=False, italic=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing: float = 1.1):
    tf = _frame(slide, x, y, w, h, anchor)
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        _run(p, line, size, color, bold, italic)
    return tf


def ftext(slide, x, y, w, h, value: str, max_pt: float, min_pt: float, color: str, bold=False, italic=False,
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing: float = 1.1) -> float:
    """Text box whose font size is fitted to the box. Returns the size used."""
    size = fit([value], w, h, max_pt, min_pt, spacing=spacing, bold=bold)
    text(slide, x, y, w, h, value, size, color, bold, italic, align, anchor, spacing)
    return size


def _label(shape, value: str, size: float, color: str) -> None:
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = 0
    tf.text = value
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = tf.paragraphs[0].runs[0]
    r.font.name, r.font.size, r.font.bold = font(), Pt(size), True
    r.font.color.rgb = rgb(color)


def _bullet(p, size: float, color: str, char: str = "●") -> None:
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Pt(size * 1.4)))
    pPr.set("indent", str(-Pt(size * 1.4)))
    clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(clr, qn("a:srgbClr")).set("val", color)
    etree.SubElement(pPr, qn("a:buSzPct")).set("val", "100000")
    etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
    etree.SubElement(pPr, qn("a:buChar")).set("char", char)


def body(slide, t, x, y, w, h, lead: str, items: list[str], max_pt: float = 20, min_pt: float = 10.5,
         accent: str | None = None, anchor=MSO_ANCHOR.TOP) -> float:
    """Lead paragraph + bullets fitted together into one box. 'Head: text' bullets get a bold head."""
    accent = accent or t["accent"]
    items = [i for i in items if i and i.strip()]
    paras = ([lead] if lead else []) + items
    size = fit(paras, w, h, max_pt, min_pt, spacing=1.1, gap=0.55, indent=1.4, first_scale=1.06 if lead else 1.0)
    tf = _frame(slide, x, y, w, h, anchor)
    first = True
    if lead:
        p = tf.paragraphs[0]
        p.line_spacing = 1.1
        p.space_after = Pt(size * 0.75)
        _run(p, lead, size * 1.06, t["muted"])
        first = False
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.line_spacing = 1.1
        p.space_after = Pt(size * 0.55)
        _bullet(p, size, accent)
        head, sep, rest = item.partition(":")
        if sep and 2 <= len(head) <= 45 and rest.strip():
            _run(p, head + ":", size, t["text"], bold=True)
            _run(p, rest, size, t["text"])
        else:
            _run(p, item, size, t["text"])
    return size


def picture(slide, path: Path, x, y, w, h, rounded: float | None = 0.05, darken: float = 1.0, blur: int = 0):
    """Crops the photo to the frame aspect first, so it is never stretched."""
    px_w = 1600
    px_h = max(1, int(px_w * h / w))
    out = path.with_name(f"{path.stem}_{px_w}x{px_h}_{int(darken * 100)}_{blur}.jpg")
    if not out.exists():
        img = fit_image(Image.open(path).convert("RGB"), px_w, px_h)
        if blur:
            img = img.filter(ImageFilter.GaussianBlur(blur))
        if darken != 1.0:
            img = ImageEnhance.Brightness(img).enhance(darken)
        img.save(out, "JPEG", quality=84, optimize=True)
    pic = slide.shapes.add_picture(str(out), I(x), I(y), I(w), I(h))
    if rounded:
        pic.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
        geom = pic._element.spPr.find(qn("a:prstGeom"))
        av = geom.find(qn("a:avLst"))
        if av is None:
            av = etree.SubElement(geom, qn("a:avLst"))
        gd = etree.SubElement(av, qn("a:gd"))
        gd.set("name", "adj")
        gd.set("fmla", f"val {int(rounded * 100000)}")
    return pic


def is_portrait(path: Path | None) -> bool:
    """Portraits and small pictures (old photos of people) are framed instead of stretched over the slide."""
    if not path:
        return False
    try:
        with Image.open(path) as im:
            w, h = im.size
    except Exception:
        return False
    return h > w * 1.02 or w < 1100


def photo_box(slide, t, img: Path, x, y, w, h) -> None:
    """A photo in the deck's style: rounded, framed, square or a circle."""
    kind = style()["photo"]
    if kind == "circle":
        d = min(w, h)
        cx, cy = x + w / 2, y + h / 2
        outline(slide, cx - d / 2 - 0.13, cy - d / 2 - 0.13, d + 0.26, d + 0.26, t["accent"], 2.5, MSO_SHAPE.OVAL)
        picture(slide, img, cx - d / 2, cy - d / 2, d, d, rounded=None).auto_shape_type = MSO_SHAPE.OVAL
    elif kind == "frame":
        outline(slide, x + 0.16, y + 0.16, w, h, t["accent"], 1.25)
        picture(slide, img, x, y, w, h, rounded=None)
    elif kind in ("square", "bleed"):
        picture(slide, img, x, y, w, h, rounded=None)
    else:
        picture(slide, img, x, y, w, h)


def transition(slide, kind: str = "fade") -> None:
    el = slide._element
    tr = etree.SubElement(el, qn("p:transition"))
    tr.set("spd", "med")
    etree.SubElement(tr, qn(f"p:{kind}"))
    # schema order: cSld, clrMapOvr, transition, timing
    after = el.find(qn("p:clrMapOvr"))
    if after is None:
        after = el.find(qn("p:cSld"))
    after.addnext(tr)


def background(slide, t: dict) -> None:
    fill = slide.background.fill
    if style()["bg"] == "gradient":
        fill.gradient()
        fill.gradient_angle = t["angle"]
        stops = fill.gradient_stops
        stops[0].color.rgb = rgb(t["bg"][0])
        stops[0].position = 0
        stops[1].color.rgb = rgb(t["bg"][1])
        stops[1].position = 1.0
    else:
        fill.solid()
        fill.fore_color.rgb = rgb(t["bg"][0])
    if style()["frame"]:
        outline(slide, 0.22, 0.22, 13.333 - 0.44, 7.5 - 0.44, t["accent"], 0.75)


def decor(slide, t: dict, seed: int) -> None:
    """Background ornament in the topic's style: gives every slide depth without clutter."""
    if not style()["decor"]:
        return
    op = 0.16 if t["dark"] else 0.11
    pattern = "dots" if style()["key"] == "playful" else t.get("pattern", "blobs")
    a, b = (t["accent"], t["accent2"]) if seed % 2 else (t["accent2"], t["accent"])
    if pattern == "geometric":
        for i, (x, y, s) in enumerate([(11.6, -1.1, 2.6), (12.5, 0.8, 1.3), (-1.0, 6.2, 2.2)]):
            sh = rect(slide, x, y, s, s, a if i % 2 else b, op, MSO_SHAPE.HEXAGON if i != 1 else MSO_SHAPE.DIAMOND)
            sh.rotation = 15 * (seed % 3)
    elif pattern == "lines":
        for i in range(4):
            rect(slide, 12.75 + i * 0.16, -0.5, 0.04, 8.5, a, op * (1.8 - i * 0.35))
        rect(slide, 0, 7.36, 13.333, 0.05, b, op * 2)
    elif pattern == "dots":
        for r in range(3):
            for c in range(5):
                rect(slide, 11.25 + c * 0.32, 0.22 + r * 0.32, 0.1, 0.1, a if (r + c) % 2 else b, op * 2.4,
                     MSO_SHAPE.OVAL)
        rect(slide, -1.4, 6.1, 2.6, 2.6, b, op * 0.8, MSO_SHAPE.OVAL)
    elif seed % 2:
        rect(slide, 11.2, -2.0, 3.8, 3.8, a, op, MSO_SHAPE.OVAL)
        rect(slide, -1.5, 6.0, 3.0, 3.0, b, op * 0.8, MSO_SHAPE.OVAL)
    else:
        rect(slide, -1.8, -2.0, 3.8, 3.8, a, op * 0.8, MSO_SHAPE.OVAL)
        rect(slide, 11.4, 5.9, 3.4, 3.4, b, op, MSO_SHAPE.OVAL)


def footer(slide, t: dict, deck_title: str, n: int, total: int, x0: float = X0, x1: float = X1) -> None:
    text(slide, x0, 6.98, max(1.0, min(8.0, x1 - x0 - 1.9)), 0.3, deck_title.upper()[:80], 9, t["muted"], bold=True)
    text(slide, x1 - 1.6, 6.98, 1.6, 0.3, f"{n:02d} / {total:02d}", 9, t["muted"], bold=True, align=PP_ALIGN.RIGHT)


def heading(slide, t: dict, value: str, x=X0, w=X1 - X0, max_pt: float = 32) -> None:
    """Slide title with the style's decoration (bars, rule, coloured band, dot or pill)."""
    kind = style()["title"]
    if kind == "band":
        left = 0.0 if x <= X0 + 0.01 else x - 0.35
        right = 13.333 if x + w >= X1 - 0.01 else x + w + 0.2
        rect(slide, left, 0.42, right - left, 1.36, t["accent"])
        ftext(slide, x, 0.5, w, 1.2, value, max_pt, 20, on(t["accent"]), bold=True, anchor=MSO_ANCHOR.MIDDLE,
              spacing=1.0)
        return
    tx, tw = (x + 0.42, w - 0.42) if kind == "dot" else (x, w)
    ftext(slide, tx, 0.5, tw, 1.12, value, max_pt, 20, t["text"], bold=True, anchor=MSO_ANCHOR.BOTTOM, spacing=1.0)
    if kind == "rule":
        rect(slide, x, 1.79, w, 0.02, t["muted"], 0.6)
        rect(slide, x, 1.72, 0.16, 0.16, t["accent"])
    elif kind == "dot":
        rect(slide, x, 1.3, 0.24, 0.24, t["accent"], shape=MSO_SHAPE.OVAL)
    elif kind == "pill":
        rect(slide, x, 1.7, 1.5, 0.15, t["accent"], shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        rect(slide, x + 1.62, 1.7, 0.15, 0.15, t["accent2"], shape=MSO_SHAPE.OVAL)
    else:
        rect(slide, x, 1.74, 0.9, 0.07, t["accent"], shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        rect(slide, x + 1.0, 1.74, 0.25, 0.07, t["accent2"], shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)


def badge(slide, t, x, y, n: int, color: str) -> None:
    """Card number in the style's look: a circle, a square or just a big number."""
    label = f"{n:02d}"
    kind = style()["badge"]
    if kind == "text":
        text(slide, x, y - 0.06, 1.3, 0.74, label, 28, color, bold=True, anchor=MSO_ANCHOR.MIDDLE, spacing=1.0)
    else:
        shape = MSO_SHAPE.RECTANGLE if kind == "square" else MSO_SHAPE.OVAL
        _label(rect(slide, x, y, 0.62, 0.62, color, shape=shape), label, 14, on(color))


def lead_line(slide, t, lead: str, y: float = TOP, max_h: float = 0.95) -> float:
    """Lead sentence under the heading across the slide. Returns the y where content can start."""
    if not lead:
        return y
    size = fit([lead], X1 - X0, max_h, 17, 12)
    h = min(max_h, n_lines(lead, X1 - X0, size) * size * LINE_H * 1.1 / 72 + 0.05)
    text(slide, X0, y, X1 - X0, h, lead, size, t["muted"])
    return y + h + 0.22


# ───────────────────────── slide layouts ─────────────────────────
# A layout may return the free horizontal range (x0, x1) for the footer when a photo takes a slide edge.

def _meta(t, author: str) -> str:
    return " · ".join(x for x in [author, str(dt.date.today().year)] if x)


def s_cover(slide, t, deck, img, author):
    kind = style()["cover"]
    if is_portrait(img):
        kind = "framed"                                 # a portrait is never stretched across the slide
    if kind == "framed":
        decor(slide, t, 1)
        photo_box(slide, t, img, 7.85, 0.85, 4.8, 5.8)
        rect(slide, X0, 1.75, 0.12, 3.6, t["accent"])
        ftext(slide, 1.2, 1.0, 6.2, 2.9, deck["title"], 46, 24, t["text"], bold=True, anchor=MSO_ANCHOR.BOTTOM,
              spacing=0.95)
        ftext(slide, 1.2, 4.1, 6.0, 1.4, deck.get("subtitle", ""), 19, 12, t["muted"])
        text(slide, 1.2, 6.2, 6, 0.4, _meta(t, author), 14, t["accent"], bold=True)
        return
    picture(slide, img, 0, 0, 13.333, 7.5, rounded=None, darken=0.55 if kind == "split" else 0.5)
    meta_color = t["accent2"] if t["dark"] else "FFFFFF"
    if kind == "center":
        rect(slide, 0, 0, 13.333, 7.5, shade(t), 0.45)
        ftext(slide, 1.3, 1.2, 10.73, 2.75, deck["title"], 50, 26, "FFFFFF", bold=True, align=PP_ALIGN.CENTER,
              anchor=MSO_ANCHOR.BOTTOM, spacing=0.95)
        rect(slide, 6.17, 4.12, 1.0, 0.08, t["accent"])
        ftext(slide, 1.8, 4.4, 9.73, 1.2, deck.get("subtitle", ""), 20, 13, "E5E7EB", align=PP_ALIGN.CENTER)
        text(slide, 1.3, 6.2, 10.73, 0.4, _meta(t, author), 14, meta_color, bold=True, align=PP_ALIGN.CENTER)
        return
    rect(slide, 0, 0, 8.6, 7.5, shade(t), 0.6)
    rect(slide, X0, 1.75, 0.12, 3.6, t["accent"])
    ftext(slide, 1.2, 1.2, 7.1, 2.75, deck["title"], 48, 26, "FFFFFF", bold=True, anchor=MSO_ANCHOR.BOTTOM,
          spacing=0.95)
    ftext(slide, 1.2, 4.15, 7.0, 1.35, deck.get("subtitle", ""), 20, 13, "E5E7EB")
    text(slide, 1.2, 6.25, 7, 0.4, _meta(t, author), 14, meta_color, bold=True)


def s_image_side(slide, t, s, img, right: bool):
    kind = style()["photo"]
    if kind == "bleed":
        iw = 5.0
        img_x = 13.333 - iw if right else 0.0
        picture(slide, img, img_x, 0, iw, 7.5, rounded=None)
        tx = X0 if right else iw + 0.6
        tw = (img_x - 0.5 - tx) if right else (X1 - tx)
    else:
        iw = 4.55
        img_x = 13.333 - 0.6 - iw if right else 0.6
        tx = X0 if right else 0.6 + iw + 0.5
        tw = (img_x - 0.45 - tx) if right else (X1 - tx)
        photo_box(slide, t, img, img_x, 0.6, iw, 6.1)
        if kind == "rounded":
            rect(slide, img_x + (0.3 if right else iw - 1.0), 6.32, 0.7, 0.7, t["accent"], shape=MSO_SHAPE.OVAL)
    heading(slide, t, s["title"], x=tx, w=tw, max_pt=30)
    body(slide, t, tx, TOP, tw, BOTTOM - TOP, s.get("lead", ""), s.get("bullets") or [], 20, 10.5)
    if kind == "bleed":
        return (X0, img_x - 0.5) if right else (tx, X1)
    return None


def s_article(slide, t, s):
    heading(slide, t, s["title"])
    lead, items = s.get("lead", ""), s.get("bullets") or []
    if lead and items:
        lw = 4.2
        rect(slide, X0, TOP + 0.05, 0.09, BOTTOM - TOP - 0.1, t["accent"])
        ftext(slide, X0 + 0.35, TOP, lw - 0.35, BOTTOM - TOP, lead, 22, 12, t["text"], anchor=MSO_ANCHOR.MIDDLE,
              spacing=1.15)
        px = X0 + lw + 0.35
    else:
        px = X0
        items = items or [lead]
    panel(slide, t, px, TOP - 0.1, X1 - px, BOTTOM - TOP + 0.2)
    body(slide, t, px + 0.4, TOP + 0.2, X1 - px - 0.8, BOTTOM - TOP - 0.4, "", items, 19, 10.5,
         anchor=MSO_ANCHOR.MIDDLE)


def s_cards(slide, t, s):
    heading(slide, t, s["title"])
    top = lead_line(slide, t, s.get("lead", ""))
    items = (s.get("bullets") or [])[:4] or [s.get("lead") or s["title"]]
    n = len(items)
    gap = 0.3
    cw = (X1 - X0 - gap * (n - 1)) / n
    ch = BOTTOM - top
    parts = []
    for item in items:
        head, sep, rest = item.partition(":")
        parts.append((head.strip(), rest.strip()) if sep and len(head) < 70 and rest.strip() else ("", item))
    heads = [h for h, _ in parts if h]
    head_h = 1.1 if heads else 0.0
    body_top = 1.12 + head_h
    head_size = min(fit([h], cw - 0.6, head_h, 19, 10, max_lines=3, bold=True, spacing=1.0) for h in heads) \
        if heads else 0
    body_size = min(fit([b], cw - 0.6, ch - body_top - 0.25, 19 if n <= 3 else 18, 9.5) for _, b in parts)
    for i, (head, rest) in enumerate(parts):
        x = X0 + (cw + gap) * i
        color = t["accent"] if i % 2 == 0 else t["accent2"]
        panel(slide, t, x, top, cw, ch, 0.07, border=color)
        if style()["card"] != "outline":
            rect(slide, x + 0.25, top, cw - 0.5, 0.07, color)
        badge(slide, t, x + 0.3, top + 0.32, i + 1, color)
        if head:
            text(slide, x + 0.3, top + 1.08, cw - 0.6, head_h, head, head_size, t["text"], bold=True,
                 anchor=MSO_ANCHOR.TOP, spacing=1.0)
        text(slide, x + 0.3, top + body_top, cw - 0.6, ch - body_top - 0.25, rest, body_size,
             t["muted"] if head else t["text"], spacing=1.1)


def s_stats(slide, t, s):
    stats = (s.get("stats") or [])[:4]
    if not stats:
        return s_cards(slide, t, s)
    heading(slide, t, s["title"])
    top = lead_line(slide, t, s.get("lead", ""))
    n = len(stats)
    gap = 0.3
    cw = (X1 - X0 - gap * (n - 1)) / n
    rest = [b for b in (s.get("bullets") or [])[:3] if b]
    sh = 2.45 if sum(len(b) for b in rest) > 220 else 2.8
    vh = 1.05
    for i, st in enumerate(stats):
        x = X0 + (cw + gap) * i
        color = t["accent"] if i % 2 == 0 else t["accent2"]
        panel(slide, t, x, top, cw, sh, 0.08, border=color if style()["card"] == "outline" else None)
        rect(slide, x + 0.3, top + sh - 0.07, cw - 0.6, 0.07, color)
        value = st["value"].replace(" ", " ")       # the number stays on one line
        size = fit([value], (cw - 0.4) / 1.12, vh, 50, 14, max_lines=1, bold=True, spacing=1.0)  # wide digits
        tf = text(slide, x + 0.2, top + 0.15, cw - 0.4, vh, value, size, color, bold=True, align=PP_ALIGN.CENTER,
                  anchor=MSO_ANCHOR.MIDDLE, spacing=1.0)
        tf.word_wrap = False
        ftext(slide, x + 0.25, top + 0.25 + vh, cw - 0.5, sh - vh - 0.45, st["label"], 16, 9.5, t["text"],
              align=PP_ALIGN.CENTER, spacing=1.05)
    if rest:
        y = top + sh + 0.35
        body(slide, t, X0, y, X1 - X0, BOTTOM - y, "", rest, 19, 10.5)


def s_quote(slide, t, s, img):
    q = s.get("quote") or s["title"]
    if is_portrait(img):
        # the speaker's portrait in a frame, the quote beside it on the theme background
        decor(slide, t, 0)
        photo_box(slide, t, img, 8.55, 0.9, 4.1, 5.6)
        text(slide, 0.9, 0.45, 2, 1.6, "“", 120, t["accent"], bold=True)
        ftext(slide, 1.5, 1.45, 6.6, 3.1, q, 30, 16, t["text"], italic=True, anchor=MSO_ANCHOR.MIDDLE, spacing=1.15)
        if s.get("quote_author"):
            rect(slide, 1.5, 4.95, 0.6, 0.06, t["accent"])
            ftext(slide, 2.3, 4.75, 5.8, 0.5, s["quote_author"], 18, 12, t["accent"], bold=True,
                  anchor=MSO_ANCHOR.MIDDLE)
        if s.get("lead"):
            ftext(slide, 1.5, 5.45, 6.6, 1.25, s["lead"], 16, 10.5, t["muted"], spacing=1.1)
        return
    picture(slide, img, 0, 0, 13.333, 7.5, rounded=None, darken=0.4)
    rect(slide, 0, 0, 13.333, 7.5, shade(t), 0.4)
    text(slide, 1.1, 0.45, 2, 1.6, "“", 140, t["accent"] if t["dark"] else t["accent2"], bold=True)
    ftext(slide, 1.8, 1.45, 9.9, 3.1, q, 34, 18, "FFFFFF", italic=True, anchor=MSO_ANCHOR.MIDDLE, spacing=1.15)
    y = 4.75
    if s.get("quote_author"):
        rect(slide, 1.8, y + 0.2, 0.6, 0.06, t["accent"])
        ftext(slide, 2.6, y, 9, 0.5, s["quote_author"], 18, 12, "FFFFFF", bold=True, anchor=MSO_ANCHOR.MIDDLE)
    if s.get("lead"):
        ftext(slide, 1.8, y + 0.75, 9.9, 1.25, s["lead"], 17, 11, "E5E7EB", spacing=1.1)


def s_two_column(slide, t, s):
    heading(slide, t, s["title"])
    top = lead_line(slide, t, s.get("lead", ""))
    cols = [(s.get("left_title") or "", s.get("bullets") or [], t["accent"]),
            (s.get("right_title") or "", s.get("right_bullets") or [], t["accent2"])]
    cw = (X1 - X0 - 0.4) / 2
    hh = 0.75
    bh = BOTTOM - top - hh - 0.45
    size = min(fit([i for i in items if i], cw - 0.7, bh, 18, 10, spacing=1.1, gap=0.55, indent=1.4)
               for _, items, _ in cols)
    for i, (title, items, color) in enumerate(cols):
        x = X0 + (cw + 0.4) * i
        panel(slide, t, x, top, cw, BOTTOM - top, 0.05, border=color if style()["card"] == "outline" else None)
        rect(slide, x, top, cw, hh, color, 1.0, band_shape(), 0.25 if band_shape() != MSO_SHAPE.RECTANGLE else None)
        ftext(slide, x + 0.35, top, cw - 0.7, hh, title, 20, 12, on(color), bold=True, anchor=MSO_ANCHOR.MIDDLE,
              spacing=1.0)
        body(slide, t, x + 0.35, top + hh + 0.25, cw - 0.7, bh, "", items, size, size, accent=color)
    vs = rect(slide, 6.665 - 0.4, top + (BOTTOM - top) / 2 - 0.4, 0.8, 0.8, t["bg"][1], shape=MSO_SHAPE.OVAL)
    vs.line.color.rgb = rgb(t["accent"])
    vs.line.width = Pt(2)
    _label(vs, "VS", 14, t["text"])


def s_timeline(slide, t, s):
    items = (s.get("bullets") or [])[:6]
    if not items:
        return s_article(slide, t, s)
    heading(slide, t, s["title"])
    top = lead_line(slide, t, s.get("lead", ""))
    n = len(items)
    step = (X1 - X0) / n
    bw = step - 0.2
    line_y = top + 0.78
    rect(slide, X0, line_y, X1 - X0, 0.05, t["muted"], 0.5)
    parts = []
    for i, item in enumerate(items):
        head, sep, rest = item.partition(":")
        parts.append((head.strip(), rest.strip()) if sep and len(head) <= 30 else (f"{i + 1}", item))
    box_top = line_y + 0.42
    box_h = BOTTOM - box_top
    size = min(fit([r], bw - 0.3, box_h - 0.3, 17 if n <= 4 else 16, 10) for _, r in parts)
    head_size = min(fit([h], bw, 0.6, 20, 10, max_lines=1, bold=True) for h, _ in parts)
    for i, (head, rest) in enumerate(parts):
        cx = X0 + step * i + step / 2
        color = t["accent"] if i % 2 == 0 else t["accent2"]
        dot = rect(slide, cx - 0.2, line_y - 0.18, 0.4, 0.4, color,
                   shape=MSO_SHAPE.RECTANGLE if style()["badge"] == "square" else MSO_SHAPE.OVAL)
        dot.line.color.rgb = rgb(t["bg"][0])
        dot.line.width = Pt(4)
        text(slide, cx - bw / 2, top, bw, 0.6, head, head_size, color, bold=True, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.BOTTOM, spacing=1.0)
        panel(slide, t, cx - bw / 2, box_top, bw, box_h, 0.08)
        text(slide, cx - bw / 2 + 0.15, box_top + 0.15, bw - 0.3, box_h - 0.3, rest, size, t["text"],
             align=PP_ALIGN.CENTER, spacing=1.1)


def s_thanks(slide, t, deck, img):
    if is_portrait(img):
        decor(slide, t, 1)
        ftext(slide, 1, 2.1, 11.33, 1.8, deck.get("thanks") or "Рахмет!", 54, 30, t["text"], bold=True,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=1.0)
        rect(slide, 6.07, 4.05, 1.2, 0.08, t["accent"], shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        ftext(slide, 1.5, 4.35, 10.33, 1.0, deck["title"], 20, 13, t["muted"], align=PP_ALIGN.CENTER)
        return
    picture(slide, img, 0, 0, 13.333, 7.5, rounded=None, darken=0.45, blur=18)
    rect(slide, 0, 0, 13.333, 7.5, shade(t), 0.35)
    ftext(slide, 1, 2.1, 11.33, 1.8, deck.get("thanks") or "Рахмет!", 54, 30, "FFFFFF", bold=True,
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=1.0)
    rect(slide, 6.07, 4.05, 1.2, 0.08, t["accent"], shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    ftext(slide, 1.5, 4.35, 10.33, 1.0, deck["title"], 20, 13, "E5E7EB", align=PP_ALIGN.CENTER)


# ───────────────────────── entry point ─────────────────────────

def pick_font(requested: str, deck: dict) -> str:
    if requested in FONTS:
        return requested
    auto = (deck.get("design") or {}).get("font")
    return auto if auto in FONTS else "Arial"


def build(deck: dict, theme: str, image_paths: dict[str, Path], out: Path, author: str = "",
          font_name: str = "auto", style_name: str = "auto") -> Path:
    """theme: palette name or "auto" (from deck["design"]); font_name: one of FONTS or "auto";
    style_name: one of STYLES or "auto". image_paths: "cover" and slide indices ("0", "1", …) → jpg files."""
    _font.set(pick_font(font_name, deck))
    _style.set(resolve_style(style_name, deck))
    t = resolve(theme, deck)
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]
    slides = deck["slides"]
    total = len(slides) + 2
    kinds = ["fade", "push", "wipe"]

    sl = prs.slides.add_slide(blank)
    background(sl, t)
    s_cover(sl, t, deck, image_paths["cover"], author)
    transition(sl, "fade")

    for i, s in enumerate(slides):
        sl = prs.slides.add_slide(blank)
        background(sl, t)
        layout = s.get("layout")
        img = image_paths.get(str(i))
        if layout in IMAGE_LAYOUTS and not img:
            layout = "article"
        if layout != "quote":
            decor(sl, t, i)
        free = None
        if layout in ("image_right", "image_left"):
            free = s_image_side(sl, t, s, img, right=(layout == "image_right"))
        elif layout == "quote":
            s_quote(sl, t, s, img)
        elif layout == "stats":
            s_stats(sl, t, s)
        elif layout == "two_column":
            s_two_column(sl, t, s)
        elif layout == "timeline":
            s_timeline(sl, t, s)
        elif layout == "cards":
            s_cards(sl, t, s)
        else:
            s_article(sl, t, s)
        if layout != "quote":
            footer(sl, t, deck["title"], i + 2, total, *(free or (X0, X1)))
        if s.get("notes"):
            sl.notes_slide.notes_text_frame.text = s["notes"]
        transition(sl, kinds[i % len(kinds)])

    sl = prs.slides.add_slide(blank)
    background(sl, t)
    s_thanks(sl, t, deck, image_paths["cover"])
    transition(sl, "fade")

    prs.core_properties.title = deck["title"]
    prs.core_properties.author = author or "ZhubaiSlide"
    prs.save(out)
    return out
