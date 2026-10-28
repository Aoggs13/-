from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Emu

from app import builder, images
from app.themes import AUTO, FONTS, THEMES, contrast, from_design, resolve
from sample_deck import DECK


def _paths(d: Path, deck: dict, colors) -> dict:
    paths = {"cover": d / "cover.jpg"}
    images.gradient_art("cover", colors).save(paths["cover"])
    for i, s in enumerate(deck["slides"]):
        if builder.needs_image(s):
            paths[str(i)] = d / f"{i}.jpg"
            images.gradient_art(str(i), colors).save(paths[str(i)])
    return paths


@pytest.mark.parametrize("theme", [AUTO, *THEMES])
def test_build_every_theme(tmp_path: Path, theme: str):
    t = resolve(theme, DECK)
    out = builder.build(DECK, theme, _paths(tmp_path, DECK, (t["bg"][0], t["accent"], t["accent2"])),
                        tmp_path / "deck.pptx", author="Test", font_name=AUTO)
    prs = Presentation(out)
    assert len(prs.slides) == len(DECK["slides"]) + 2
    # every shape stays on the slide
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                assert sh.left >= 0 and sh.top >= 0
                assert sh.left + sh.width <= prs.slide_width + Emu(10)


@pytest.mark.parametrize("font", list(FONTS))
def test_fonts_are_applied(tmp_path: Path, font: str):
    out = builder.build(DECK, "aurora", _paths(tmp_path, DECK, ("000000", "FFFFFF", "888888")),
                        tmp_path / "deck.pptx", font_name=font)
    names = {r.font.name for s in Presentation(out).slides for sh in s.shapes if sh.has_text_frame
             for p in sh.text_frame.paragraphs for r in p.runs}
    assert names == {font}


def test_auto_font_comes_from_design():
    assert builder.pick_font(AUTO, DECK) == "Times New Roman"
    assert builder.pick_font("Calibri", DECK) == "Calibri"
    assert builder.pick_font(AUTO, {"design": {"font": "Comic Sans"}}) == "Arial"


def test_fitted_text_fits_its_box():
    paras = [s for s in DECK["slides"][1]["bullets"]]
    size = builder.fit(paras, 6.5, 4.0, 20, 9, spacing=1.1, gap=0.55, indent=1.4)
    assert builder.text_height(paras, 6.5, size, spacing=1.1, gap=0.55, indent=1.4) <= 4.0
    assert size < 20                                    # long text really got smaller
    assert builder.fit(["Қысқа"], 6.5, 4.0, 20, 9) == 20  # short text keeps the big size


@pytest.mark.parametrize("design", [
    {"dark": True, "bg1": "FFFFFF", "bg2": "EEEEEE", "accent": "FFFF00", "accent2": "00FFFF"},  # wrong: light bg
    {"dark": False, "bg1": "000000", "bg2": "222222", "accent": "FFFFFF", "accent2": "EEEEEE"},  # wrong: dark bg
    {"dark": True, "bg1": "nonsense", "bg2": "", "accent": "#12", "accent2": "zz"},              # garbage
])
def test_auto_palette_is_always_readable(design):
    t = from_design(design)
    for bg in t["bg"]:
        assert contrast(t["text"], bg) >= 4.5
        assert contrast(t["accent"], bg) >= 2.5


def test_fit_image_keeps_aspect():
    from PIL import Image
    img = Image.new("RGB", (3000, 1000))
    assert images.fit(img, 800, 600).size == (800, 600)
